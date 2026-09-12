#!/usr/bin/env python3
"""Build a human-readable Excel workbook from processing data."""
from __future__ import annotations
import argparse,csv
from pathlib import Path
from collections import Counter,defaultdict
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

def read(path):
    with Path(path).open(encoding="utf-8-sig",newline="") as f:return list(csv.DictReader(f))
def add_sheet(wb,name,headers,rows):
    ws=wb.create_sheet(name); ws.append(headers)
    for r in rows: ws.append([r.get(h,"") if isinstance(r,dict) else r for h in headers])
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        width=min(max(len(str(x.value or "")) for x in col)+2,45); ws.column_dimensions[get_column_letter(col[0].column)].width=width
    return ws

def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--state-dir",type=Path,required=True); ap.add_argument("--output",type=Path,required=True); a=ap.parse_args()
    logs=read(a.state_dir/"processing_log.csv"); obs=read(a.state_dir/"observations.csv")
    latest={}
    for r in logs: latest[r["work_item_id"]]=r
    statuses=Counter(r["status"] for r in latest.values()); rules=Counter(r["rule_id"] for r in latest.values() if r["status"]=="COMPLETE")
    rule_obs=Counter(r["rule_id"] for r in obs); classes=Counter(r["information_class"] for r in obs); residual=Counter(r["source_element"] for r in obs if r.get("residual"))
    runs=Counter(r["run_id"] for r in logs)
    wb=Workbook(); wb.remove(wb.active)
    summary=[
      {"Metric":"Work items (current state)","Value":len(latest)},
      {"Metric":"Complete","Value":statuses["COMPLETE"]},{"Metric":"Failed","Value":statuses["FAILED"]},
      {"Metric":"Observations","Value":len(obs)},{"Metric":"CORE observations","Value":classes["CORE"]},
      {"Metric":"ASSOCIATED observations","Value":classes["ASSOCIATED"]},{"Metric":"RESIDUAL observations","Value":classes["RESIDUAL"]},
      {"Metric":"Runs represented in log","Value":len(runs)},
    ]
    add_sheet(wb,"Run Summary",["Metric","Value"],summary)
    add_sheet(wb,"Processing Status",list(next(iter(latest.values())).keys()) if latest else ["work_item_id"],list(latest.values()))
    add_sheet(wb,"Rule Statistics",["rule_id","completed_work_items","observations"],[{"rule_id":k,"completed_work_items":v,"observations":rule_obs[k]} for k,v in sorted(rules.items())])
    add_sheet(wb,"Information Classes",["information_class","observations"],[{"information_class":k,"observations":v} for k,v in sorted(classes.items())])
    add_sheet(wb,"Review Queue",["work_item_id","source_id","source_record","rule_id","status","exception"],[r for r in latest.values() if r["status"]!="COMPLETE"])
    add_sheet(wb,"Observations",list(obs[0].keys()) if obs else ["run_id"],obs)
    add_sheet(wb,"Residual",["source_element","count"],[{"source_element":k,"count":v} for k,v in residual.most_common()])
    a.output.parent.mkdir(parents=True,exist_ok=True); wb.save(a.output); print(a.output)
if __name__=="__main__":main()
