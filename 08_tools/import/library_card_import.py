"""MusicStewardship — D-side Library Card exchange importer prototype.

Rev A public encoding:
* CSV is the Library Card data artifact.
* 0..n and 1..n fields are compact JSON arrays in one CSV cell.
* 1 fields are scalar CSV values.
* D records its own processing revision.

The workbook test package remains supported as a convenience wrapper.
"""
from __future__ import annotations
import argparse,csv,json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "04_src"))

from library_card_interface import certified_interface
from typing import Any
D_PROCESSING_REVISION="library_card_import.v0.1"
def read_csv_table(path:Path):
    with path.open("r",newline="",encoding="utf-8-sig") as f:return list(csv.DictReader(f))
def rows_from_workbook(path:Path,sheet:str):
    import openpyxl; wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    if sheet not in wb.sheetnames: raise ValueError(f"Required sheet not found: {sheet}")
    return [list(r) for r in wb[sheet].iter_rows(values_only=True)]
def table_from_workbook(path:Path,sheet:str):
    data=rows_from_workbook(path,sheet)
    if not data: raise ValueError(f"{sheet} is empty")
    headers=["" if x is None else str(x) for x in data[0]]; out=[]
    for row in data[1:]:
        if not any(x not in (None,"") for x in row): continue
        out.append({headers[i]:"" if i>=len(row) or row[i] is None else str(row[i]) for i in range(len(headers))})
    return out
def package_info(path:Path):
    result={}
    for row in rows_from_workbook(path,"Package Info"):
        if len(row)>=2 and row[0] is not None: result[str(row[0])]="" if row[1] is None else str(row[1])
    return result
def contract_for_card(card):
    revision=card.get("library_card_interface_revision","")
    if not revision:
        raise ValueError(
            "Library Card does not declare library_card_interface_revision."
        )
    return certified_interface(revision)

def decode_public_value(field,value,contract):
    if value=="":
        return [] if field in contract.array_fields else ""
    if field in contract.array_fields:
        try:
            decoded=json.loads(value)
        except json.JSONDecodeError as exc:
            raise ValueError(
                f"Field {field!r} must contain a JSON array in "
                f"interface revision {contract.revision}."
            ) from exc
        if not isinstance(decoded,list):
            raise ValueError(
                f"Field {field!r} must contain a JSON array in "
                f"interface revision {contract.revision}."
            )
        return decoded
    return value

def validate_card_shape(card,contract):
    expected=set(contract.field_names)
    actual=set(card)
    missing=sorted(expected-actual)
    extra=sorted(actual-expected)

    if missing or extra:
        details=[]
        if missing:
            details.append(f"missing certified fields: {missing}")
        if extra:
            details.append(f"undeclared fields: {extra}")
        raise ValueError(
            f"Library Card does not match certified "
            f"{contract.interface_id} revision {contract.revision}: "
            + "; ".join(details)
        )

def decode_card(card):
    contract=contract_for_card(card)
    validate_card_shape(card,contract)
    return {
        field:decode_public_value(field,card[field],contract)
        for field in contract.field_names
    }

def verify_standard(card):
    contract=contract_for_card(card)
    result=[]
    for definition in contract.fields:
        field=definition["Field"]
        cardinality=definition["Cardinality"]
        meaning=definition["Meaning"]
        required=cardinality in ("1","1..n")

        if field not in card:
            result.append(
                (field,"FAIL","Field absent from Library Card header.")
            )
            continue

        if required and card[field]=="":
            result.append(
                (field,"REVIEW","Required field is present but blank.")
            )
            continue

        try:
            decoded=decode_public_value(field,card[field],contract)
        except ValueError as exc:
            result.append((field,"FAIL",str(exc)))
            continue

        if cardinality in ("0..n","1..n") and not isinstance(decoded,list):
            result.append(
                (field,"FAIL","Multi-valued field is not encoded as a JSON array.")
            )
        else:
            result.append((field,"PASS",f"Represented: {meaning}"))

    return result
def import_csv(cards_csv):
    cards=read_csv_table(cards_csv)
    if not cards: raise ValueError("No Library Card records found.")
    imported=[]; checks=[]
    for card in cards:
        imported_card=decode_card(card); imported_card["d_processing_revision"]=D_PROCESSING_REVISION; imported_card["d_import_interface_revision"]=card.get("library_card_interface_revision",""); imported.append(imported_card)
        checks += verify_standard(card)
    return cards,imported,checks
def import_workbook(workbook):
    info=package_info(workbook)
    certified_interface(info.get("Interface revision",""))
    cards=table_from_workbook(workbook,"Library Cards")
    if not cards: raise ValueError("No Library Card records found.")
    imported=[]; checks=[]
    for card in cards:
        imported_card=decode_card(card); imported_card["d_processing_revision"]=D_PROCESSING_REVISION; imported_card["d_import_interface_revision"]=info["Interface revision"]; imported.append(imported_card)
        checks += verify_standard(card)
    return cards,imported,checks
def write_outputs(outdir,interface_revision,imported,checks):
    outdir.mkdir(parents=True,exist_ok=True)
    (outdir/"D_imported_library_cards.json").write_text(json.dumps({"schema":"musicstewardship.d_imported_library_card.v0.1","d_processing_revision":D_PROCESSING_REVISION,"library_card_interface_revision":interface_revision,"records":imported},indent=2,ensure_ascii=False),encoding="utf-8")
    with (outdir/"D_verification.csv").open("w",newline="",encoding="utf-8") as f:
        w=csv.writer(f); w.writerow(["check","result","detail"]); w.writerows(checks)
    print(f"Interface revision   : {interface_revision}"); print(f"D processing revision: {D_PROCESSING_REVISION}"); print(f"Cards imported       : {len(imported)}"); print(f"Verification checks  : {len(checks)}"); print(f"PASS                 : {sum(r[1]=='PASS' for r in checks)}"); print(f"REVIEW               : {sum(r[1]=='REVIEW' for r in checks)}"); print(f"FAIL                 : {sum(r[1]=='FAIL' for r in checks)}"); print(f"Output               : {outdir}")
def main():
    ap=argparse.ArgumentParser(); source=ap.add_mutually_exclusive_group(required=True); source.add_argument("--csv",type=Path); source.add_argument("--workbook",type=Path); ap.add_argument("--output-dir",type=Path,required=True); args=ap.parse_args()
    if args.csv: _,imported,checks=import_csv(args.csv); interface_revision=imported[0].get("library_card_interface_revision","") if imported else ""
    else: info=package_info(args.workbook); _,imported,checks=import_workbook(args.workbook); interface_revision=info.get("Interface revision","")
    certified_interface(interface_revision)
    write_outputs(args.output_dir,interface_revision,imported,checks)
if __name__=="__main__": main()
